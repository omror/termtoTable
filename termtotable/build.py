# -*- coding: utf-8 -*-
"""Olayları normalize tabloya çevirip CSV ve XLSX yazar."""

import csv
from datetime import date
from pathlib import Path

from holidays import aralikta_tatiller

GUNLER = ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma", "Cumartesi", "Pazar"]

KOLONLAR = ["id", "ogretim_yili", "kategori", "olay", "baslangic", "bitis",
            "baslangic_gunu", "bitis_gunu", "takvim_gunu", "okul_tatili",
            "kademe", "dogrulama", "kaynak"]

BASLIK = ["ID", "Öğretim Yılı", "Kategori", "Olay", "Başlangıç", "Bitiş",
          "Başlangıç Günü", "Bitiş Günü", "Takvim Günü", "Okul Tatili mi?",
          "Kademe", "Doğrulama", "Kaynak"]

RENK = {
    "Dönem": "DDEBF7", "Ara Tatil": "FFF2CC", "Yarıyıl Tatili": "FFF2CC",
    "Resmî Tatil": "E2EFDA", "Dinî Bayram": "E2EFDA",
    "Mesleki Çalışma": "F2F2F2", "Uyum Eğitimi": "F2F2F2", "Yıl Sonu": "DDEBF7",
}

ETIKET = {
    "mesleki_calisma": ("Mesleki Çalışma", "Öğretmenlerin yıl başı mesleki çalışmaları", "Hayır", "Öğretmen"),
    "uyum_egitimi":    ("Uyum Eğitimi", "Okul öncesi ve ilkokul 1. sınıf uyum eğitimi", "Hayır", "Okul öncesi + İlkokul 1"),
    "birinci_donem":   ("Dönem", "1. dönem (birinci yarıyıl)", "Hayır", "Tüm kademeler"),
    "ara_tatil_1":     ("Ara Tatil", "1. dönem ara tatili", "Evet", "Tüm kademeler"),
    "yariyil_tatili":  ("Yarıyıl Tatili", "Yarıyıl (sömestr) tatili", "Evet", "Tüm kademeler"),
    "ikinci_donem":    ("Dönem", "2. dönem (ikinci yarıyıl)", "Hayır", "Tüm kademeler"),
    "ara_tatil_2":     ("Ara Tatil", "2. dönem ara tatili", "Evet", "Tüm kademeler"),
    "yil_sonu":        ("Yıl Sonu", "Eğitim-öğretim yılının sona ermesi / karne", "Hayır", "Tüm kademeler"),
}


def tablo_kur(olaylar, ogretim_yili, kaynak_url):
    satirlar = []

    for kod, (b, e) in olaylar.items():
        kategori, ad, tatil, kademe = ETIKET.get(kod, ("Diğer", kod, "?", "?"))
        satirlar.append({
            "kategori": kategori, "olay": ad,
            "baslangic": b, "bitis": e,
            "okul_tatili": tatil, "kademe": kademe,
            "dogrulama": "meb", "kaynak": kaynak_url,
        })

    if "birinci_donem" in olaylar and "ikinci_donem" in olaylar:
        yil_basi = olaylar["birinci_donem"][0]
        yil_sonu = olaylar["ikinci_donem"][1]
        if yil_sonu:
            for t in aralikta_tatiller(yil_basi, yil_sonu):
                satirlar.append({
                    "kategori": t["kategori"], "olay": t["ad"],
                    "baslangic": t["baslangic"], "bitis": t["bitis"],
                    "okul_tatili": "Evet", "kademe": "Tüm kademeler",
                    "dogrulama": t["dogrulama"],
                    "kaynak": ("2429 sayılı Ulusal Bayram ve Genel Tatiller Kanunu"
                               if t["kategori"] == "Resmî Tatil"
                               else "Diyanet İşleri Başkanlığı dinî günler takvimi"),
                })

    satirlar.sort(key=lambda r: (r["baslangic"], r["olay"]))
    for i, s in enumerate(satirlar, 1):
        b, e = s["baslangic"], s["bitis"]
        s["id"] = i
        s["ogretim_yili"] = ogretim_yili
        s["baslangic_gunu"] = GUNLER[b.weekday()]
        s["bitis_gunu"] = GUNLER[e.weekday()] if e else ""
        s["takvim_gunu"] = (e - b).days + 1 if e else ""
        s["baslangic"] = b.isoformat()
        s["bitis"] = e.isoformat() if e else ""

    return satirlar


def csv_yaz(satirlar, yol):
    Path(yol).parent.mkdir(parents=True, exist_ok=True)
    with open(yol, "w", newline="", encoding="utf-8-sig") as f:
        yazici = csv.DictWriter(f, fieldnames=KOLONLAR, extrasaction="ignore")
        yazici.writeheader()
        yazici.writerows(satirlar)


def xlsx_yaz(satirlar, yol, kaynak_url):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    Path(yol).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Takvim"

    ince = Side(style="thin", color="BFBFBF")
    kenar = Border(left=ince, right=ince, top=ince, bottom=ince)

    # Başlık satırı
    for c, baslik in enumerate(BASLIK, 1):
        h = ws.cell(1, c, baslik)
        h.font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
        h.fill = PatternFill("solid", fgColor="44546A")
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        h.border = kenar

    # Veri satırları
    for r, s in enumerate(satirlar, 2):
        degerler = [
            s["id"], s["ogretim_yili"], s["kategori"], s["olay"],
            date.fromisoformat(s["baslangic"]),
            date.fromisoformat(s["bitis"]) if s["bitis"] else None,
            s["baslangic_gunu"], s["bitis_gunu"],
            # Sabit sayı değil formül: tarih değişirse gün sayısı da değişir
            '=IF(F%d="","",F%d-E%d+1)' % (r, r, r),
            s["okul_tatili"], s["kademe"], s["dogrulama"], s["kaynak"],
        ]
        for c, v in enumerate(degerler, 1):
            hucre = ws.cell(r, c, v)
            hucre.font = Font(name="Arial", size=10)
            hucre.border = kenar
            hucre.alignment = Alignment(
                vertical="top",
                wrap_text=c in (4, 11, 13),
                horizontal="center" if c in (1, 5, 6, 9, 10, 12) else "left",
            )
            if c in (5, 6):
                hucre.number_format = "DD.MM.YYYY"
            dolgu = RENK.get(s["kategori"])
            if dolgu:
                hucre.fill = PatternFill("solid", fgColor=dolgu)

    # Sütun genişlikleri, dondurma, filtre
    for i, g in enumerate([5, 12, 16, 42, 12, 12, 14, 13, 12, 13, 24, 13, 46], 1):
        ws.column_dimensions[get_column_letter(i)].width = g
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(BASLIK)), len(satirlar) + 1)

    # Alt not
    n = len(satirlar) + 3
    ws.cell(n, 1, "Kaynak ve notlar").font = Font(name="Arial", bold=True, size=10)
    notlar = [
        "MEB çalışma takvimi: %s" % kaynak_url,
        "Resmî tatiller: 2429 sayılı Ulusal Bayram ve Genel Tatiller Kanunu",
        "Doğrulama: meb = genelgeden okundu · kanun = kanundan hesaplandı ·",
        "   diyanet = Diyanet takvimiyle teyitli · hesaplanan = hicri dönüşüm, ±1 gün sapabilir",
        "'Takvim Günü' hafta sonları dahildir, iş günü sayısı değildir.",
        "Üretim tarihi: %s" % date.today().strftime("%d.%m.%Y"),
    ]
    for j, t in enumerate(notlar, 1):
        ws.cell(n + j, 1, "• " + t).font = Font(name="Arial", size=9, color="595959")

    wb.save(yol)


if __name__ == "__main__":
    from parser import ayikla
    from validate import dogrula

    metin = (
        "Okullarda birinci dönem, 14 Eylül 2026 Pazartesi günü başlayacak ve "
        "22 Ocak 2027 Cuma günü sona erecek. "
        "2026-2027 eğitim ve öğretim yılında birinci dönem ara tatili 16 Kasım 2026 "
        "Pazartesi günü başlayacak ve 20 Kasım 2026 Cuma günü sona erecek. "
        "Yarıyıl tatili 25 Ocak 2027 Pazartesi günü başlayıp 5 Şubat 2027 Cuma günü "
        "tamamlanacak. İkinci dönem ise 8 Şubat 2027 Pazartesi günü başlayacak. "
        "İkinci dönemin ara tatili 8 Mart 2027 Pazartesi günü başlayacak ve "
        "12 Mart 2027 Cuma günü sona erecek."
    )
    url = "https://www.meb.gov.tr/2026-2027-egitim-ogretim-yili-takvimi-aciklandi/haber/41057/tr"

    olaylar, uyarilar = ayikla(metin)
    olaylar["ikinci_donem"] = (olaylar["ikinci_donem"][0], date(2027, 6, 25))
    dogrula(olaylar, "2026-2027")

    satirlar = tablo_kur(olaylar, "2026-2027", url)
    csv_yaz(satirlar, "out/termtotable_2026-2027.csv")
    xlsx_yaz(satirlar, "out/termtotable_2026-2027.xlsx", url)

    print("%d kayıt yazıldı:" % len(satirlar))
    print("  out/termtotable_2026-2027.csv")
    print("  out/termtotable_2026-2027.xlsx")